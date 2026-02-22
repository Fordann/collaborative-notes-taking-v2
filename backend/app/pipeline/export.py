import subprocess
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from odf.opendocument import load as odf_load
from odf.text import P, Span
from odf.style import Style, TextProperties, ParagraphProperties


def _style_docx_additions(docx_path: str) -> None:
    """Post-process a .docx file to style [AJOUT] paragraphs.

    Spec section 11:
    - Paragraph background: green pâle (#E8F5E9)
    - [AJOUT] prefix: bold, color #2E7D32
    - Source note (italic text at end): gray
    """
    doc = Document(docx_path)

    for para in doc.paragraphs:
        full_text = para.text
        if "[AJOUT]" not in full_text:
            continue

        # Apply light green background shading to the paragraph
        pPr = para._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "E8F5E9")
        pPr.append(shd)

        # Rebuild runs to style [AJOUT] prefix in bold green
        # and source notes in italic gray
        original_text = full_text
        para.clear()

        # Split around [AJOUT]
        before, _, after = original_text.partition("[AJOUT]")

        if before.strip():
            run_before = para.add_run(before)
            run_before.font.size = Pt(11)

        # [AJOUT] prefix in bold green
        run_tag = para.add_run("[AJOUT]")
        run_tag.bold = True
        run_tag.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
        run_tag.font.size = Pt(11)

        # Detect source note at the end: (Source : ...) — Pandoc strips the
        # markdown italic markers *, so we look for "(Source" directly.
        remaining = after
        source_text = ""
        source_start = remaining.rfind("(Source")
        if source_start != -1 and remaining.rstrip().endswith(")"):
            source_text = remaining[source_start:]
            remaining = remaining[:source_start]

        # Main content
        if remaining:
            run_content = para.add_run(remaining)
            run_content.font.size = Pt(11)

        # Source in italic gray
        if source_text:
            run_source = para.add_run(source_text)
            run_source.italic = True
            run_source.font.color.rgb = RGBColor(0x75, 0x75, 0x75)
            run_source.font.size = Pt(10)

    doc.save(docx_path)


def _odf_get_text(element) -> str:
    """Recursively extract plain text from an ODF element."""
    text = ""
    for node in element.childNodes:
        if hasattr(node, "data"):
            text += node.data
        else:
            text += _odf_get_text(node)
    return text


def _style_odt_additions(odt_path: str) -> None:
    """Post-process an .odt file to style [AJOUT] paragraphs.

    Same visual logic as docx via ODF styles.
    """
    doc = odf_load(odt_path)

    # Create styles for additions
    ajout_para_style = Style(name="AjoutParagraph", family="paragraph")
    ajout_para_style.addElement(ParagraphProperties(backgroundcolor="#E8F5E9"))
    doc.automaticstyles.addElement(ajout_para_style)

    ajout_tag_style = Style(name="AjoutTag", family="text")
    ajout_tag_style.addElement(TextProperties(
        fontweight="bold",
        color="#2E7D32",
        fontsize="11pt",
    ))
    doc.automaticstyles.addElement(ajout_tag_style)

    ajout_source_style = Style(name="AjoutSource", family="text")
    ajout_source_style.addElement(TextProperties(
        fontstyle="italic",
        color="#757575",
        fontsize="10pt",
    ))
    doc.automaticstyles.addElement(ajout_source_style)

    # Collect paragraphs first to avoid modifying tree during iteration
    paragraphs = list(doc.body.getElementsByType(P))

    for elem in paragraphs:
        text_content = _odf_get_text(elem)

        if "[AJOUT]" not in text_content:
            continue

        # Build a replacement paragraph instead of clearing children in-place,
        # because removeChild() fails on text nodes (AssertionError in odfpy).
        before, _, after = text_content.partition("[AJOUT]")

        new_para = P(stylename=ajout_para_style)

        if before.strip():
            new_para.addText(before)

        # [AJOUT] tag in bold green
        tag_span = Span(stylename=ajout_tag_style)
        tag_span.addText("[AJOUT]")
        new_para.addElement(tag_span)

        # Detect source note — Pandoc strips markdown italic markers
        remaining = after
        source_text = ""
        source_start = remaining.rfind("(Source")
        if source_start != -1 and remaining.rstrip().endswith(")"):
            source_text = remaining[source_start:]
            remaining = remaining[:source_start]

        if remaining:
            new_para.addText(remaining)

        if source_text:
            source_span = Span(stylename=ajout_source_style)
            source_span.addText(source_text)
            new_para.addElement(source_span)

        # Replace old paragraph with new styled one
        parent = elem.parentNode
        parent.insertBefore(new_para, elem)
        parent.removeChild(elem)

    doc.save(odt_path)


async def export_to_original_format(
    enriched_md: str,
    original_file_path: str,
    original_format: str,
) -> str:
    """Reconvertit le Markdown enrichi dans le format original."""

    original = Path(original_file_path)
    output_dir = original.parent / "output"
    output_dir.mkdir(exist_ok=True)

    md_path = output_dir / "enriched.md"
    md_path.write_text(enriched_md, encoding="utf-8")
    output_path = output_dir / f"enriched{original.suffix}"

    if original_format == "md":
        output_path.write_text(enriched_md, encoding="utf-8")

    elif original_format == "txt":
        subprocess.run(
            ["pandoc", str(md_path), "-t", "plain", "-o", str(output_path)],
            capture_output=True, text=True, timeout=30, check=True,
        )

    elif original_format == "docx":
        subprocess.run(
            ["pandoc", str(md_path), f"--reference-doc={original_file_path}", "-o", str(output_path)],
            capture_output=True, text=True, timeout=30, check=True,
        )
        _style_docx_additions(str(output_path))

    elif original_format == "odt":
        subprocess.run(
            ["pandoc", str(md_path), "-o", str(output_path)],
            capture_output=True, text=True, timeout=30, check=True,
        )
        _style_odt_additions(str(output_path))

    return str(output_path)
